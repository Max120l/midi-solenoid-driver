#!/usr/bin/env python3
"""
book_to_organ -- turn a read book (row-indexed events) into an organ-format score.

Input: the events.json from read_book_video.py, and the organ's layout sheet
with a `book scale` column giving, for every solenoid row, the book key
number that drives it. Output: the multi-track organ-format MIDI the arranger
takes (tracks named as in the sheet, real pitches from the sheet), so

    book_to_organ.py valse-brune.events.json --scale limonaire49t_scale.xlsx --top-key 49 -o valse-brune.fororgan.mid
    organ_arranger.py valse-brune.fororgan.mid --organ instrument/organ.yaml

Rows are numbered from the top of the picture. Which key the top row is
depends on the scanning software: --top-key 49 means the top row is key 49
and keys descend down the picture; --top-key 1 the reverse. Find out by
listening, or with the audio correlation notes in the README.

No transposition happens here: the book's key N plays whatever the sheet
says key N plays on this organ.
"""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

import mido
import openpyxl

__version__ = "0.1.0"

TRACK_CHANNEL = {"Melody": 0, "Main": 0, "TenorCM": 1, "TrebCM": 2, "Accompainment": 3, "Drums": 0, "Registers": 0}


def read_scale(path: Path) -> dict[int, tuple[str, int, str]]:
    """book key -> (track, written note, label)."""
    ws = openpyxl.load_workbook(path, data_only=True).active
    header = [c.value for c in ws[1]]
    if "book scale" not in [str(h).strip().lower() if h else "" for h in header]:
        raise SystemExit("the sheet needs a 'book scale' column")
    bcol = [str(h).strip().lower() if h else "" for h in header].index("book scale")
    # track groups: a header cell names the track for itself and the unnamed cells after it
    groups = []
    for i, h in enumerate(header):
        if h and i != 0 and i != bcol:
            groups.append((i, str(h).strip()))
    keymap: dict[int, tuple[str, int, str]] = {}
    for r in ws.iter_rows(min_row=2, values_only=True):
        key = r[bcol]
        if key is None or key == "":
            continue
        for gi, (start, track) in enumerate(groups):
            end = groups[gi + 1][0] if gi + 1 < len(groups) else bcol
            cells = r[start:end]
            if cells and cells[0] not in (None, ""):
                note = int(cells[0])
                label = str(cells[1]) if len(cells) > 1 and cells[1] not in (None, "") else ""
                keymap[int(key)] = (track, note, label)
                break
    return keymap


def convert(events: dict[int, list[list[float]]], keymap: dict[int, tuple[str, int, str]], keys: int,
            top_key: int, title: str) -> tuple[mido.MidiFile, dict]:
    tpb, tempo = 480, 500_000
    mid = mido.MidiFile(type=1, ticks_per_beat=tpb)
    cond = mido.MidiTrack()
    cond.append(mido.MetaMessage("track_name", name=title, time=0))
    cond.append(mido.MetaMessage("set_tempo", tempo=tempo, time=0))
    cond.append(mido.MetaMessage("end_of_track", time=0))
    mid.tracks.append(cond)
    per_track: dict[str, list[tuple[int, int, int]]] = {}
    report = {"placed": 0, "unmapped_rows": [], "per_track": {}}
    for row, ivs in events.items():
        key = top_key - row if top_key == keys else top_key + row
        if key not in keymap:
            if ivs:
                report["unmapped_rows"].append((row, key, len(ivs)))
            continue
        track, note, _ = keymap[key]
        for s, e in ivs:
            a = int(round(mido.second2tick(s, tpb, tempo)))
            b = max(int(round(mido.second2tick(e, tpb, tempo))), a + 1)
            per_track.setdefault(track, []).append((a, 1, note))
            per_track[track].append((b, 0, note))
            report["placed"] += 1
    order = [t for t in TRACK_CHANNEL if t in per_track] + [t for t in per_track if t not in TRACK_CHANNEL]
    for track in order:
        evs = sorted(per_track[track])
        tr = mido.MidiTrack()
        tr.append(mido.MetaMessage("track_name", name=track, time=0))
        prev = 0
        ch = TRACK_CHANNEL.get(track, 0)
        for tick, on, note in evs:
            tr.append(mido.Message("note_on" if on else "note_off", channel=ch, note=note,
                                   velocity=100 if on else 0, time=tick - prev))
            prev = tick
        tr.append(mido.MetaMessage("end_of_track", time=0))
        mid.tracks.append(tr)
        report["per_track"][track] = len(evs) // 2
    return mid, report


def main(argv: list[str] | None = None) -> int:
    p = argparse.ArgumentParser(prog="book_to_organ", description="Row events from a read book -> organ-format MIDI.")
    p.add_argument("events", help="events.json from read_book_video")
    p.add_argument("--scale", required=True, help="layout sheet with a 'book scale' column")
    p.add_argument("--top-key", type=int, required=True, help="book key number of the top row (1 or the key count)")
    p.add_argument("--keys", type=int, default=49, help="keys on the book (default 49)")
    p.add_argument("-o", "--output", help="organ-format .mid to write (default: EVENTS.fororgan.mid)")
    p.add_argument("--title", default=None)
    p.add_argument("--version", action="version", version=f"%(prog)s {__version__}")
    a = p.parse_args(argv)
    src = Path(a.events)
    events = {int(k): v for k, v in json.load(open(src)).items()}
    keymap = read_scale(Path(a.scale))
    if a.top_key not in (1, a.keys):
        p.error("--top-key must be 1 or the key count")
    out = Path(a.output) if a.output else src.with_name(src.name.replace(".events.json", "") + ".fororgan.mid")
    mid, report = convert(events, keymap, a.keys, a.top_key, a.title or src.stem.replace(".events", ""))
    mid.save(str(out))
    print(f"wrote {out}: {report['placed']} notes; per track {report['per_track']}")
    for row, key, n in report["unmapped_rows"]:
        print(f"  warning: row {row} (key {key}) has {n} holes but no entry in the sheet")
    return 0


if __name__ == "__main__":
    sys.exit(main())
